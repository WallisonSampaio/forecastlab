"""Downloadable Excel upload template (Portuguese sheets, headers, instructions and examples).

Every Portuguese header is a column alias of ``schema.map_columns`` (data sheet) or ``schema.map_plan_columns``
(plan sheet), and ``readers.read_table`` picks the "Dados" sheet, so a filled-in template uploads as is.
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.exports.excel_style import SUB_FONT, TITLE_FONT, signature_rows, style_table

from .schema import (
    DESCRIPTIONS,
    DISPLAY,
    HEADERS_PT,
    OPTIONAL,
    PLAN_DESCRIPTIONS,
    PLAN_HEADERS_PT,
    PLAN_REQUIRED,
    REQUIRED,
)

SHEET_DATA = "Dados"
SHEET_INSTRUCTIONS = "Instruções"
SHEET_PLAN = "Plano (opcional)"

# Canonical column -> Portuguese header used in the template (kept as a module name for existing imports).
HEADERS = HEADERS_PT

# English names of the plan columns, also accepted on upload.
PLAN_DISPLAY = {
    "date": "Date",
    "customer": "Customer",
    "sku": "SKU",
    "budget_units": "BudgetUnits",
    "target_units": "TargetUnits",
    "budget_revenue": "BudgetRevenue",
    "target_revenue": "TargetRevenue",
}

_ATLAS = {
    "region": "Sudeste",
    "channel": "Varejo Moderno",
    "customer": "Atlas Hipermercados",
    "sku": "SKU-10101",
    "customer_location": "São Paulo/SP",
    "customer_potential": 250000,
    "product": "AquaViva Água Mineral 500ml",
    "category": "Bebidas",
    "brand": "AquaViva",
}
_CLICKCESTA = {
    "region": "Sudeste",
    "channel": "E-commerce",
    "customer": "ClickCesta Online",
    "sku": "SKU-20101",
    "customer_location": "São Paulo/SP - CD",
    "customer_potential": 90000,
    "product": "Crocantes Batata Clássica 90g",
    "category": "Snacks",
    "brand": "Crocantes",
}


def _row(series: dict, month: int, units: int, price: float, **drivers: float) -> dict:
    return {
        **series,
        "date": datetime(2024, month, 1),
        "units": units,
        "revenue": round(units * price, 2),
        "price": price,
        "promotion": drivers.get("promotion", 0),
        "distribution": drivers["distribution"],
        "stockout": drivers.get("stockout", 0),
        "holiday": drivers["holiday"],
        "marketing_investment": drivers["marketing"],
        "merchandising_event": drivers.get("merchandising", 0),
    }


# Six months of one series (the minimum history accepted for a forecast) plus a second series, so the
# example rows themselves pass the validation.
EXAMPLE = [
    _row(_ATLAS, 1, 1250, 2.49, distribution=78.5, holiday=1, marketing=450.0),
    _row(_ATLAS, 2, 1410, 2.34, promotion=1, distribution=79.0, holiday=2, marketing=900.0, merchandising=1),
    _row(_ATLAS, 3, 1180, 2.49, distribution=79.2, stockout=1, holiday=1, marketing=450.0),
    _row(_ATLAS, 4, 1225, 2.49, distribution=79.4, holiday=1, marketing=450.0),
    _row(_ATLAS, 5, 1290, 2.49, distribution=79.8, holiday=2, marketing=500.0),
    _row(_ATLAS, 6, 1180, 2.49, distribution=80.1, holiday=0, marketing=450.0),
    _row(_CLICKCESTA, 5, 320, 6.99, distribution=55.0, holiday=2, marketing=120.0),
    _row(_CLICKCESTA, 6, 345, 6.99, distribution=55.5, holiday=0, marketing=120.0),
]

# Plan sheet: Portuguese headers (aliases of schema.map_plan_columns) and one example row.
PLAN_COLUMNS = ["date", "customer", "sku", "budget_units", "target_units", "budget_revenue", "target_revenue"]
PLAN_HEADERS = [PLAN_HEADERS_PT[c] for c in PLAN_COLUMNS]
PLAN_EXAMPLE = [datetime(2025, 1, 1), "Atlas Hipermercados", "SKU-10101", 1300, 1400, 3237.0, 3486.0]
DATE_FORMAT = "dd/mm/yyyy"

INSTRUCTIONS = [
    "Preencha a aba 'Dados' com uma linha por Cliente x SKU x Mês. Linhas diárias ou semanais são agregadas por "
    "mês automaticamente.",
    "Apague as linhas de exemplo antes de colar os seus dados. São necessários pelo menos 6 meses de histórico; "
    "com 24 meses ou mais os modelos capturam a sazonalidade.",
    "Formatos aceitos: .xlsx (recomendado), .csv e .parquet. Os nomes das colunas são reconhecidos sem diferenciar "
    "maiúsculas, minúsculas e acentos, em português ou em inglês (veja a coluna 'Nome alternativo').",
    f"A aba '{SHEET_PLAN}' recebe orçamento e meta por Cliente x SKU x Mês. Inclua todas as combinações "
    "Cliente x SKU da aba 'Dados' (as que ficarem de fora não terão orçamento nem meta) e os últimos 12 meses do "
    "histórico, para comparar o realizado com o orçamento. Sem essa aba, o plano é derivado do histórico: "
    "orçamento = mesmo mês do ano anterior + 5%; meta = orçamento + 5%.",
]


def _column_table(ws, title: str, rows: list[list[str]]) -> None:
    ws.append([])
    ws.append([title])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    ws.append(["Coluna", "Obrigatória", "Descrição", "Nome alternativo"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
        ws.cell(row=ws.max_row, column=3).alignment = Alignment(wrap_text=True, vertical="top")


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_DATA
    cols = REQUIRED + OPTIONAL
    ws.append([HEADERS[c] for c in cols])
    for row in EXAMPLE:
        ws.append([row[c] for c in cols])
    style_table(
        ws,
        {
            HEADERS["date"]: DATE_FORMAT,
            HEADERS["units"]: "units",
            HEADERS["customer_potential"]: "units",
            HEADERS["revenue"]: "currency2",
            HEADERS["price"]: "currency2",
            HEADERS["marketing_investment"]: "currency2",
        },
    )

    info = wb.create_sheet(SHEET_INSTRUCTIONS)
    info["A1"] = "ForecastLab — Modelo de planilha para carga de dados"
    info["A1"].font = TITLE_FONT
    for i, text in enumerate(INSTRUCTIONS, start=2):
        cell = info.cell(row=i, column=1, value=text)
        cell.font = SUB_FONT
        # One paragraph per row across the 4 table columns, wrapped (no text running past column D).
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        info.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        info.row_dimensions[i].height = 15 * (len(text) // 140 + 1)
    _column_table(
        info,
        f"Aba '{SHEET_DATA}'",
        [[HEADERS[c], "Sim" if c in REQUIRED else "Opcional", DESCRIPTIONS[c], DISPLAY[c]] for c in cols],
    )
    _column_table(
        info,
        f"Aba '{SHEET_PLAN}'",
        [
            [PLAN_HEADERS_PT[c], "Sim" if c in PLAN_REQUIRED else "Opcional", PLAN_DESCRIPTIONS[c], PLAN_DISPLAY[c]]
            for c in PLAN_COLUMNS
        ],
    )
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 13
    info.column_dimensions["C"].width = 100
    info.column_dimensions["D"].width = 22
    info.sheet_view.showGridLines = False
    signature_rows(info, info.max_row + 3)

    plan = wb.create_sheet(SHEET_PLAN)
    plan.append(PLAN_HEADERS)
    plan.append(PLAN_EXAMPLE)
    style_table(
        plan,
        {
            PLAN_HEADERS_PT["date"]: DATE_FORMAT,
            PLAN_HEADERS_PT["budget_units"]: "units",
            PLAN_HEADERS_PT["target_units"]: "units",
            PLAN_HEADERS_PT["budget_revenue"]: "currency",
            PLAN_HEADERS_PT["target_revenue"]: "currency",
        },
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
