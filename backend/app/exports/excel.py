"""Excel exports: full planning workbook and executive summary (Portuguese sheets and headers).

Internal codes (series types, models, classifications ...) are translated to their Portuguese
labels on the way out; the numbers keep native Excel formats, which a pt-BR Excel renders as
1.234.567 / R$ 1.234 / 14,8%.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

from app.alerts.engine import build_alerts
from app.config import APP_NAME, APP_VERSION
from app.fva.fva import fva_report
from app.scenarios.engine import ScenarioParams, simulate_rows
from app.scenarios.store import compare, get_scenario
from app.services import ptbr
from app.services.context import PlanningContext, filter_ids
from app.services.forecast_views import leaderboard, performance_table
from app.sop.cockpit import cockpit
from app.sop.feasibility import feasibility_by

from .excel_style import FORMATS, NAVY, SUB_FONT, TITLE_FONT, signature_rows, style_table

DISCLAIMER = (
    "Todos os dados da base de demonstração são sintéticos e fictícios. Os forecasts são estimativas "
    "estatísticas; valide-os no processo de S&OP antes de comprometer planos de abastecimento ou financeiros."
)
NO_DATA = "Sem dados para esta seleção."
ALL = "Todos"

# Sheet names (Excel limit: 31 characters).
SHEET_COVER = "Capa"
SHEET_SUMMARY = "Resumo Executivo"
SHEET_DETAIL = "Detalhe do Forecast"
SHEET_LEADERBOARD = "Ranking de Modelos"
SHEET_PERFORMANCE = "Performance das Séries"
SHEET_SCENARIOS = "Cenários"
SHEET_FEASIBILITY = "Viabilidade da Meta"
SHEET_FVA = "FVA"
SHEET_ALERTS = "Alertas"
SHEET_OVERRIDES = "Ajustes Comerciais"

# API field -> Portuguese column header.
COLUMN_LABELS = {
    "series_id": "Série",
    "key": "Série",
    "date": "Mês",
    "start": "Início",
    "quarter": "Trimestre",
    "customer": "Cliente",
    "sku": "SKU",
    "product": "Produto",
    "region": "Região",
    "channel": "Canal",
    "category": "Categoria",
    "brand": "Marca",
    "series_type": "Tipo de série",
    "champion": "Modelo campeão",
    "challenger": "Modelo desafiante",
    "confidence": "Confiança",
    "confidence_score": "Score de confiança",
    "history_months": "Meses de histórico",
    "rank": "Posição",
    "model": "Modelo",
    "family": "Família",
    "wmape": "WMAPE",
    "accuracy": "Acurácia",
    "bias": "Viés",
    "bias_class": "Classe de viés",
    "rmse": "RMSE",
    "challenger_wmape": "WMAPE do desafiante",
    "gap_vs_challenger": "Vantagem vs desafiante",
    "wmape_std": "Desvio do WMAPE entre folds",
    "series_evaluated": "Séries avaliadas",
    "champion_count": "Vezes campeão",
    "challenger_count": "Vezes desafiante",
    "win_rate": "Taxa de vitória",
    "scenario_id": "ID do cenário",
    "name": "Cenário",
    "kind": "Tipo",
    "units": "Unidades",
    "revenue": "Receita",
    "vs_base_units": "vs Base (unidades)",
    "vs_base_revenue": "vs Base (receita)",
    "vs_budget_units": "vs Orçamento (unidades)",
    "vs_budget_revenue": "vs Orçamento (receita)",
    "vs_target_units": "vs Meta (unidades)",
    "vs_target_revenue": "vs Meta (receita)",
    "budget_units": "Orçamento (unidades)",
    "budget_revenue": "Orçamento (receita)",
    "target_units": "Meta (unidades)",
    "target_revenue": "Meta (receita)",
    "dimension": "Dimensão",
    "member": "Membro",
    "target": "Meta",
    "budget": "Orçamento",
    "base": "Forecast estatístico",
    "statistical": "Forecast estatístico",
    "consensus": "Consenso",
    "p10": "P10",
    "p90": "P90",
    "probability": "Probabilidade de atingir a meta",
    "classification": "Classificação",
    "gap": "Gap vs orçamento",
    "gap_pct": "Gap vs orçamento (%)",
    "gap_vs_budget": "Gap vs orçamento",
    "gap_vs_target": "Gap vs meta",
    "gap_vs_target_pct": "Gap vs meta (%)",
    "required_growth": "Crescimento necessário",
    "planner": "Planejador",
    "event_type": "Tipo de evento",
    "comment": "Comentário",
    "adjustment_pct": "Ajuste (%)",
    "actual": "Real",
    "error_statistical": "Erro do forecast estatístico",
    "error_consensus": "Erro do consenso",
    "fva_units": "FVA (unidades)",
    "outcome": "Resultado",
    "direction": "Direção",
    "severity": "Severidade",
    "title": "Alerta",
    "issue": "Problema",
    "magnitude": "Magnitude",
    "potential_impact": "Impacto potencial (R$)",
    "explanation": "Explicação",
    "recommendation": "Recomendação",
    "override_id": "ID do ajuste",
    "uplift_pct": "Ajuste (%)",
    "created_by": "Responsável",
    "created_at": "Criado em",
}

# API field -> code-to-label map applied to its values.
VALUE_LABELS: dict[str, dict[str, str]] = {
    "series_type": ptbr.SERIES_TYPE,
    "champion": ptbr.MODEL,
    "challenger": ptbr.MODEL,
    "model": ptbr.MODEL,
    "family": ptbr.FAMILY,
    "confidence": ptbr.CONFIDENCE,
    "bias_class": ptbr.BIAS_CLASS,
    "classification": ptbr.CLASSIFICATION,
    "severity": ptbr.SEVERITY,
    "kind": ptbr.SCENARIO_KIND,
    "outcome": ptbr.OUTCOME,
    "direction": ptbr.DIRECTION,
    "event_type": ptbr.EVENT_TYPE,
    "planner": ptbr.PLANNER,
    "created_by": ptbr.PLANNER,
}
DATE_COLUMNS = {"date": "date", "start": "date", "created_at": "datetime"}


def _localize(df: pd.DataFrame, columns: list[str] | None = None) -> pd.DataFrame:
    """Select columns, translate coded values and rename the headers to Portuguese."""
    if columns:
        df = df[[c for c in columns if c in df.columns]]
    df = df.copy()
    for col, mapping in VALUE_LABELS.items():
        if col in df.columns:
            df[col] = df[col].map(lambda v, m=mapping: ptbr.label(m, v))
    for col in DATE_COLUMNS:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce")
    return df.rename(columns={c: COLUMN_LABELS.get(c, c) for c in df.columns})


def _formats(by_field: dict[str, str]) -> dict[str, str]:
    out = {COLUMN_LABELS.get(k, k): v for k, v in by_field.items()}
    out.update({COLUMN_LABELS[k]: v for k, v in DATE_COLUMNS.items()})
    return out


def _sheet(
    wb: Workbook,
    title: str,
    df: pd.DataFrame,
    formats: dict[str, str] | None = None,
    columns: list[str] | None = None,
    localize: bool = True,
) -> None:
    ws = wb.create_sheet(title[:31])
    if df.empty:
        ws["A1"] = NO_DATA
        return
    out = _localize(df, columns) if localize else df
    out = out.astype(object).where(out.notna(), None)
    for row in dataframe_to_rows(out, index=False, header=True):
        ws.append(row)
    style_table(ws, _formats(formats or {}) if localize else formats)


def _filters_text(filters: dict[str, Any] | None) -> str:
    parts = []
    for k, v in (filters or {}).items():
        if v:
            values = ", ".join(str(x) for x in v) if isinstance(v, list | tuple) else str(v)
            parts.append(f"{ptbr.DIMENSION.get(k, k)}: {values}")
    return "; ".join(parts) or ALL


def _cover(wb: Workbook, ctx: PlanningContext, subtitle: str, filters: dict[str, Any] | None) -> None:
    ws = wb.active
    ws.title = SHEET_COVER
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70
    ws["A1"] = APP_NAME
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ptbr.APP_TAGLINE
    ws["A2"].font = SUB_FONT
    ws["A4"] = subtitle
    ws["A4"].font = Font(bold=True, size=14, color=NAVY)
    rows = [
        ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Versão da base", ctx.version),
        ("Execução do forecast", ctx.run_id or "-"),
        ("Último mês com dados reais", ptbr.month_label(ctx.last_actual, long=True)),
        ("Horizonte do forecast", f"{len(ctx.future_dates)} meses"),
        ("Filtros", _filters_text(filters)),
        ("Versão da aplicação", APP_VERSION),
    ]
    for i, (k, v) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, color="334155")
        ws.cell(row=i, column=2, value=v)
    r = signature_rows(ws, len(rows) + 8)
    ws.cell(row=r + 1, column=1, value=DISCLAIMER).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=2)
    ws.row_dimensions[r + 1].height = 45
    for row in ws.iter_rows(min_row=1, max_row=3):
        for c in row:
            c.fill = PatternFill(None)


def forecast_detail(
    ctx: PlanningContext, filters: dict[str, Any] | None, scenario_id: str | None = None, history_months: int = 12
) -> pd.DataFrame:
    """Series-month detail (actuals + forecast + plan) with Portuguese headers and labels."""
    ids = filter_ids(ctx, filters)
    s = ctx.series[ctx.series["series_id"].isin(ids)]
    attrs = s[
        [
            "series_id",
            "region",
            "channel",
            "customer",
            "sku",
            "product",
            "category",
            "brand",
            "series_type",
            "champion",
            "challenger",
            "wmape",
            "bias",
            "confidence",
        ]
    ]
    f = ctx.forecast[ctx.forecast["series_id"].isin(ids)].copy()
    f["row_type"] = "Forecast"
    if scenario_id:
        sc = get_scenario(scenario_id)
        rows = simulate_rows(ctx, ScenarioParams.from_dict(sc["params"]))
        f = f.merge(
            rows[["series_id", "date", "scenario_units", "scenario_revenue"]], on=["series_id", "date"], how="left"
        )
    hist_start = ctx.last_actual - pd.DateOffset(months=history_months - 1)
    a = ctx.actuals[ctx.actuals["series_id"].isin(ids) & (ctx.actuals["date"] >= hist_start)][
        ["series_id", "date", "units", "revenue"]
    ].assign(row_type="Real")
    df = pd.concat([a, f], ignore_index=True)
    if not ctx.plan.empty:
        df = df.merge(ctx.plan, on=["series_id", "date"], how="left")
    df = df.merge(attrs, on="series_id", how="left")
    df = df.sort_values(["customer", "sku", "date"]).reset_index(drop=True)
    out = pd.DataFrame(
        {
            "Mês": pd.to_datetime(df["date"]),
            "Tipo": df["row_type"],
            "Região": df["region"],
            "Canal": df["channel"],
            "Cliente": df["customer"],
            "SKU": df["sku"],
            "Produto": df["product"],
            "Categoria": df["category"],
            "Marca": df["brand"],
            "Unidades reais": df.get("units"),
            "Receita real": df.get("revenue"),
            "Forecast estatístico": df.get("p50"),
            "Incremental (eventos)": df.get("incremental"),
            "Forecast de consenso": df.get("consensus"),
            "Receita de consenso": df.get("consensus") * df.get("price") if "price" in df else None,
            "P10": df.get("p10"),
            "P50": df.get("p50"),
            "P90": df.get("p90"),
            "Orçamento (unidades)": df.get("budget_units"),
            "Meta (unidades)": df.get("target_units"),
            "Orçamento (receita)": df.get("budget_revenue"),
            "Meta (receita)": df.get("target_revenue"),
        }
    )
    if scenario_id:
        out["Cenário (unidades)"] = df.get("scenario_units")
        out["Cenário (receita)"] = df.get("scenario_revenue")
    out["Tipo de série"] = df["series_type"].map(lambda v: ptbr.label(ptbr.SERIES_TYPE, v))
    out["Modelo campeão"] = df["champion"].map(lambda v: ptbr.label(ptbr.MODEL, v))
    out["Modelo desafiante"] = df["challenger"].map(lambda v: ptbr.label(ptbr.MODEL, v))
    out["WMAPE"] = df["wmape"]
    out["Viés"] = df["bias"]
    out["Confiança"] = df["confidence"].map(lambda v: ptbr.label(ptbr.CONFIDENCE, v))
    return out


DETAIL_FORMATS = {
    "Mês": "date",
    "Unidades reais": "units",
    "Receita real": "currency",
    "Forecast estatístico": "units",
    "Incremental (eventos)": "units",
    "Forecast de consenso": "units",
    "Receita de consenso": "currency",
    "P10": "units",
    "P50": "units",
    "P90": "units",
    "Orçamento (unidades)": "units",
    "Meta (unidades)": "units",
    "Orçamento (receita)": "currency",
    "Meta (receita)": "currency",
    "Cenário (unidades)": "units",
    "Cenário (receita)": "currency",
    "WMAPE": "percent",
    "Viés": "percent",
}


def _kpi_rows(ctx: PlanningContext, ck: dict[str, Any]) -> list[tuple[str, Any, str | None]]:
    """(label, value, format) of the executive KPIs, revenue measure."""
    k = ck.get("kpis", {})
    feas = ck.get("feasibility") or {}
    n = len(ctx.future_dates)
    return [
        ("Real - últimos 12 meses (R$)", k.get("actual_ltm"), "currency"),
        ("Crescimento vs ano anterior", k.get("actual_growth"), "percent"),
        (f"Forecast estatístico - próximos {n} meses (R$)", k.get("statistical"), "currency"),
        ("Incremental de eventos comerciais (R$)", k.get("incremental"), "currency"),
        (f"Forecast de consenso - próximos {n} meses (R$)", k.get("consensus"), "currency"),
        ("P10 do forecast (R$)", k.get("p10"), "currency"),
        ("P90 do forecast (R$)", k.get("p90"), "currency"),
        (f"Orçamento - próximos {n} meses (R$)", k.get("budget"), "currency"),
        (f"Meta - próximos {n} meses (R$)", k.get("target"), "currency"),
        ("Gap vs orçamento (R$)", k.get("gap_vs_budget"), "currency"),
        ("Gap vs orçamento (%)", k.get("gap_vs_budget_pct"), "percent"),
        ("Gap vs meta (R$)", k.get("gap_vs_target"), "currency"),
        ("Gap vs meta (%)", k.get("gap_vs_target_pct"), "percent"),
        ("Receita em risco (R$)", k.get("risk"), "currency"),
        ("Oportunidade (R$)", k.get("opportunity"), "currency"),
        ("Crescimento do forecast vs últimos 12 meses", k.get("forecast_growth"), "percent"),
        ("Acurácia do forecast", k.get("accuracy"), "percent"),
        ("WMAPE", k.get("wmape"), "percent"),
        ("Viés", k.get("bias"), "percent"),
        ("Classe de viés", ptbr.label(ptbr.BIAS_CLASS, k.get("bias_class")), None),
        ("Viabilidade da meta", ptbr.label(ptbr.CLASSIFICATION, feas.get("classification")), None),
        ("Probabilidade de atingir a meta", feas.get("probability"), "percent"),
    ]


def _kpi_sheet(wb: Workbook, title: str, rows: list[tuple[str, Any, str | None]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(["Indicador", "Valor"])
    for label, value, _ in rows:
        ws.append([label, value if value is not None else ptbr.NA])
    style_table(ws, widths={"Indicador": 48, "Valor": 24})
    for i, (_, value, fmt) in enumerate(rows, start=2):
        if fmt and value is not None:
            ws.cell(row=i, column=2).number_format = FORMATS[fmt]


def build_workbook(
    ctx: PlanningContext, filters: dict[str, Any] | None = None, scenario_id: str | None = None
) -> bytes:
    wb = Workbook()
    _cover(wb, ctx, "Plano de Demanda e Exportação de S&OP", filters)
    ck = cockpit(ctx, filters, "revenue")
    _kpi_sheet(wb, SHEET_SUMMARY, _kpi_rows(ctx, ck))

    _sheet(wb, SHEET_DETAIL, forecast_detail(ctx, filters, scenario_id), DETAIL_FORMATS, localize=False)
    _sheet(
        wb,
        SHEET_LEADERBOARD,
        pd.DataFrame(leaderboard(ctx, filters)),
        {"wmape": "percent", "accuracy": "percent", "bias": "percent", "rmse": "number", "win_rate": "percent"},
        columns=[
            "rank",
            "model",
            "family",
            "wmape",
            "accuracy",
            "bias",
            "rmse",
            "series_evaluated",
            "champion_count",
            "challenger_count",
            "win_rate",
        ],
    )
    _sheet(
        wb,
        SHEET_PERFORMANCE,
        pd.DataFrame(performance_table(ctx, filters)),
        {
            "wmape": "percent",
            "accuracy": "percent",
            "bias": "percent",
            "challenger_wmape": "percent",
            "gap_vs_challenger": "percent",
            "wmape_std": "percent",
        },
    )
    cmp = compare(ctx, None, filters, "revenue")
    _sheet(
        wb,
        SHEET_SCENARIOS,
        pd.DataFrame(cmp["rows"]),
        {
            "units": "units",
            "revenue": "currency",
            "vs_base_units": "percent",
            "vs_base_revenue": "percent",
            "vs_budget_units": "percent",
            "vs_budget_revenue": "percent",
            "vs_target_units": "percent",
            "vs_target_revenue": "percent",
            "budget_units": "units",
            "budget_revenue": "currency",
            "target_units": "units",
            "target_revenue": "currency",
        },
        columns=[
            "name",
            "kind",
            "units",
            "revenue",
            "vs_base_units",
            "vs_base_revenue",
            "vs_budget_units",
            "vs_budget_revenue",
            "vs_target_units",
            "vs_target_revenue",
            "budget_units",
            "budget_revenue",
            "target_units",
            "target_revenue",
            "scenario_id",
        ],
    )
    feas = [
        {"dimension": ptbr.DIMENSION[dim], **row}
        for dim in ("region", "channel")
        for row in feasibility_by(ctx, dim, filters)
    ]
    _sheet(
        wb,
        SHEET_FEASIBILITY,
        pd.DataFrame(feas),
        {
            "target": "currency",
            "budget": "currency",
            "consensus": "currency",
            "p10": "currency",
            "p90": "currency",
            "probability": "percent",
            "gap_vs_target": "currency",
            "gap_vs_target_pct": "percent",
            "required_growth": "percent",
        },
    )
    fva = fva_report(ctx, filters)
    _sheet(
        wb,
        SHEET_FVA,
        pd.DataFrame(fva.get("overrides", [])),
        {
            "adjustment_pct": "percent",
            "statistical": "units",
            "consensus": "units",
            "actual": "units",
            "error_statistical": "units",
            "error_consensus": "units",
            "fva_units": "units",
        },
    )
    alerts = pd.DataFrame(
        [
            {
                **{
                    k2: a[k2]
                    for k2 in (
                        "severity",
                        "title",
                        "issue",
                        "magnitude",
                        "potential_impact",
                        "explanation",
                        "recommendation",
                    )
                },
                "customer": a["entity"].get("customer"),
                "sku": a["entity"].get("sku"),
            }
            for a in build_alerts(ctx)
        ]
    )
    _sheet(wb, SHEET_ALERTS, alerts, {"potential_impact": "currency"})
    ov = ctx.overrides.drop(columns=["dataset_version"], errors="ignore").copy()
    for col in ("customer", "sku"):
        if col in ov.columns:
            ov[col] = ov[col].where(ov[col].notna() & (ov[col] != ""), ALL)
    _sheet(
        wb,
        SHEET_OVERRIDES,
        ov,
        {"uplift_pct": "percent"},
        columns=[
            "override_id",
            "date",
            "customer",
            "sku",
            "event_type",
            "uplift_pct",
            "comment",
            "created_by",
            "created_at",
        ],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_executive_workbook(ctx: PlanningContext, filters: dict[str, Any] | None = None) -> bytes:
    wb = Workbook()
    _cover(wb, ctx, "Resumo Executivo de S&OP", filters)
    ck = cockpit(ctx, filters, "revenue")
    _kpi_sheet(wb, "Indicadores", _kpi_rows(ctx, ck))
    _sheet(
        wb,
        "Trimestres",
        pd.DataFrame(ck.get("quarters", [])),
        {
            "base": "currency",
            "consensus": "currency",
            "budget": "currency",
            "target": "currency",
            "gap_vs_budget": "currency",
        },
        columns=["quarter", "start", "base", "consensus", "budget", "target", "gap_vs_budget"],
    )
    gap_cols = [
        "customer",
        "sku",
        "product",
        "region",
        "channel",
        "category",
        "consensus",
        "budget",
        "gap",
        "gap_pct",
        "series_type",
        "champion",
        "confidence",
    ]
    gap_formats = {"consensus": "currency", "budget": "currency", "gap": "currency", "gap_pct": "percent"}
    _sheet(wb, "Principais Riscos", pd.DataFrame(ck.get("risks", [])), gap_formats, columns=gap_cols)
    _sheet(wb, "Principais Oportunidades", pd.DataFrame(ck.get("opportunities", [])), gap_formats, columns=gap_cols)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
