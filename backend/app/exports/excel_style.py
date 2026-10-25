"""Shared openpyxl styling for ForecastLab workbooks."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.config import AUTHOR

NAVY = "2D3277"
INDIGO = "3483FA"  # action blue
SLATE = "64748B"
LIGHT = "F1F5F9"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
TITLE_FONT = Font(bold=True, size=20, color=NAVY, name="Calibri")
SUB_FONT = Font(size=11, color=SLATE, name="Calibri")
LINK_FONT = Font(size=10, color=INDIGO, underline="single", name="Calibri")
THIN = Border(bottom=Side(style="thin", color="E2E8F0"))

FORMATS = {
    "units": "#,##0",
    "currency": '"R$" #,##0',
    "currency2": '"R$" #,##0.00',
    "percent": "0.0%",
    "number": "#,##0.00",
    "date": "yyyy-mm",
}


def style_table(
    ws: Worksheet, formats: dict[str, str] | None = None, start_row: int = 1, widths: dict[str, int] | None = None
) -> None:
    """Header styling, number formats by column name, widths and freeze panes."""
    formats = formats or {}
    headers = [c.value for c in ws[start_row]]
    for cell in ws[start_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[start_row].height = 30
    for idx, name in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        fmt = formats.get(str(name))
        if fmt:
            for row in ws.iter_rows(min_row=start_row + 1, min_col=idx, max_col=idx):
                for cell in row:
                    cell.number_format = FORMATS.get(fmt, fmt)
        width = (widths or {}).get(str(name)) or min(42, max(11, len(str(name or "")) + 4))
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False


def signature_rows(ws: Worksheet, row: int) -> int:
    ws.cell(row=row, column=1, value=AUTHOR["signature_en"]).font = Font(bold=True, size=11, color=NAVY)
    ws.cell(row=row + 1, column=1, value=AUTHOR["title_en"]).font = SUB_FONT
    c = ws.cell(row=row + 2, column=1, value=AUTHOR["linkedin"])
    c.hyperlink = AUTHOR["linkedin"]
    c.font = LINK_FONT
    return row + 3
